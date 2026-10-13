"""Synthetic workbooks and database snapshots. No real person's data."""
from __future__ import annotations

import datetime as dt
import uuid

import openpyxl
from openpyxl.utils import column_index_from_string

SHEETS = ['سبت ', 'الاحد ', 'الاثنين ', 'الثلاثاء ', 'الاربعاء ', 'الخميس ', 'الجمعة ']
SHEET_DAY = {'سبت': 6, 'الاحد': 0, 'الاثنين': 1, 'الثلاثاء': 2, 'الاربعاء': 3, 'الخميس': 4, 'الجمعة': 5}


def twelve_hour(minute: int) -> dt.time:
    """The workbook writes 12-hour times with no AM/PM: 13:00 is stored as 01:00."""
    h = (minute // 60) % 12
    return dt.time(h, minute % 60)


def standard_headers(start=8 * 60, end=24 * 60 + 30):
    return [twelve_hour(m) for m in range(start, end + 1, 30)]


def make_workbook(path: str, content: dict | None = None, headers: dict | None = None, extra_sheets=()):
    """content: {sheet_title_stripped: [(teacher, {cell_col: text | (text, span)})]}"""
    content = content or {}
    headers = headers or {}
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title in list(SHEETS) + list(extra_sheets):
        ws = wb.create_sheet(title)
        ws.cell(1, 1, 'اسم المعلم ')
        for i, h in enumerate(headers.get(title.strip(), standard_headers())):
            ws.cell(1, 3 + i, h)
        for r, (teacher, cells) in enumerate(content.get(title.strip(), []), start=3):
            ws.cell(r, 1, teacher)
            for col, val in cells.items():
                c = column_index_from_string(col)
                text, span = (val if isinstance(val, tuple) else (val, 1))
                ws.cell(r, c, text)
                if span > 1:
                    ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c + span - 1)
    wb.save(path)
    return path


def col_for(minute: int, start=8 * 60) -> str:
    """Column letter of a minute on the standard header (column C = start)."""
    return openpyxl.utils.get_column_letter(3 + (minute - start) // 30)


def uid() -> str:
    return str(uuid.uuid4())


class SnapshotBuilder:
    def __init__(self):
        self.teachers, self.students, self.lessons, self.shift_assignments = [], [], [], []
        self.template_id = uid()

    def teacher(self, name, roster=False, deleted=False):
        t = {'id': uid(), 'full_name': name, 'teacher_type': 'shift' if roster else 'hourly', 'status': 'active',
             'is_deleted': deleted, 'created_at': '2026-01-01'}
        self.teachers.append(t)
        if roster:
            for d in range(7):
                self.shift_assignments.append({'id': uid(), 'teacher_id': t['id'], 'shift_template_id': self.template_id,
                                               'day_of_week': d, 'is_active': True, 'name': 'Full-time',
                                               'start_minute': 720, 'end_minute': 1140, 't_active': True})
        return t['id']

    def student(self, name, deleted=False):
        s = {'id': uid(), 'full_name': name, 'status': 'active', 'is_deleted': deleted, 'created_at': '2026-01-01'}
        self.students.append(s)
        return s['id']

    def lesson(self, teacher_id, day, start, dur, student_ids, status='active', source=(None, None)):
        l = {'id': uid(), 'teacher_id': teacher_id, 'day_of_week': day, 'start_minute': start, 'duration_minutes': dur,
             'end_minute': start + dur, 'lifecycle_status': status, 'source_sheet': source[0], 'source_cell': source[1],
             'parts': [{'id': uid(), 'student_id': s} for s in student_ids]}
        self.lessons.append(l)
        return l['id']

    def snapshot(self) -> dict:
        return {'teachers': self.teachers, 'students': self.students, 'lessons': self.lessons,
                'shift_assignments': self.shift_assignments}


def teacher_map(entries, roster_ids=()):
    return {'schema_version': 1, 'shift_roster_teacher_ids': sorted(roster_ids), 'teachers': entries}


def confirmed(spellings, tid, basis='SINGLE_PLAUSIBLE_RECORD', candidates=None):
    return {'sheet_spellings': spellings, 'status': 'CONFIRMED', 'teacher_id': tid, 'basis': basis,
            'candidates': candidates or [tid], 'evidence': 'test'}


def unresolved(spellings, candidates, decision='decide'):
    return {'sheet_spellings': spellings, 'status': 'UNRESOLVED', 'teacher_id': None, 'candidates': candidates,
            'decision_required': decision}
