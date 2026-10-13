"""Human-readable outputs of a plan.

reconciliation.xlsx carries student names and must stay out of version
control (scripts/import-schedule/dry-run/ is git-ignored). summary.md and
teacher-classification.md carry no student names.
"""
from __future__ import annotations

from collections import Counter, defaultdict

import openpyxl
from openpyxl.styles import Font, PatternFill

from .identity import ws
from .reconcile import DECISIONS

DAY_NAMES = ['Sunday', 'Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday', 'Saturday']


def hm(minute):
    if minute is None:
        return ''
    suffix = ' (+1 day)' if minute >= 1440 else ''
    return f'{(minute // 60) % 24:02d}:{minute % 60:02d}{suffix}'


def _sheet(wb, title, header, rows, widths):
    s = wb.create_sheet(title)
    s.append(header)
    for c in s[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill('solid', fgColor='DDDDDD')
    for r in rows:
        s.append(r)
    s.freeze_panes = 'A2'
    s.auto_filter.ref = s.dimensions
    for i, w in enumerate(widths, 1):
        s.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w


def teacher_classification(parsed, tmap, snap, evidence, outcomes) -> list[dict]:
    cells = Counter(ws(c.teacher_raw) for c in parsed.lessons)
    status = defaultdict(Counter)
    for o in outcomes:
        status[ws(o.teacher_raw)][o.status] += 1
    rows = []
    roster = snap.roster_teacher_ids()
    groups = {a['teacher_id']: a['name'] for a in snap.shift_assignments if a['is_active'] and a['t_active']}
    for e in tmap.data['teachers']:
        res = tmap.resolve(e['sheet_spellings'][0])
        tid = res.teacher_id
        ev = Counter()
        for sp in e['sheet_spellings']:
            ev.update(evidence.get(ws(sp), {}))
        rows.append({
            'sheet_spellings': e['sheet_spellings'],
            'lesson_cells': sum(cells[ws(sp)] for sp in e['sheet_spellings']),
            'status': res.status,
            'teacher_id': tid,
            'teacher_name': snap.teachers[tid]['full_name'].strip() if tid else None,
            'classification': ({'shift': f'shift-based/part-time ({groups.get(tid, "?")})', 'hourly': 'hourly'}[res.classification]
                               if res.classification else 'unclassified (identity unresolved)'),
            'basis': res.basis or '',
            'candidates': [f'{snap.teachers[c]["full_name"].strip()} ({c[:8]})' for c in e.get('candidates', []) if c in snap.teachers],
            'evidence_matches': {snap.teachers[k]['full_name'].strip() + f' ({k[:8]})': v for k, v in ev.most_common(3) if k in snap.teachers},
            'note': e.get('evidence') or e.get('decision_required') or '',
            'outcomes': dict(sum((status[ws(sp)] for sp in e['sheet_spellings']), Counter())),
            'on_roster': tid in roster if tid else False,
        })
    blank = cells.get('', 0)
    if blank:
        rows.append({'sheet_spellings': ['(blank teacher cell)'], 'lesson_cells': blank, 'status': 'UNRESOLVED', 'teacher_id': None,
                     'teacher_name': None, 'classification': 'unclassified (identity unresolved)', 'basis': '', 'candidates': [],
                     'evidence_matches': {}, 'note': DECISIONS['NO_TEACHER_IN_ROW'], 'outcomes': dict(status['']), 'on_roster': False})
    return rows


def write_xlsx(path, parsed, snap, result, classification, items):
    outcomes = result['outcomes']
    st = snap.students
    te = snap.teachers
    names = lambda ids: ' + '.join(st[i]['full_name'] for i in ids) if ids else ''
    wb = openpyxl.Workbook()
    summary = wb.active
    summary.title = 'Summary'
    counts = Counter(o.status for o in outcomes)
    for row in [
        ('Workbook SHA-256', parsed.sha256), ('Source lesson cells', len(outcomes)),
        ('ALREADY_PRESENT', counts['ALREADY_PRESENT']), ('IMPORT', counts['IMPORT']),
        ('PENDING', counts['PENDING']), ('REJECTED', counts['REJECTED']),
        ('Non-lesson cells (not counted above)', len(parsed.non_lessons)),
        ('Lessons to insert', len(items)), ('Participant rows to insert', sum(len(i['student_ids']) for i in items)),
    ]:
        summary.append(list(row))
    summary.column_dimensions['A'].width = 40
    summary.column_dimensions['B'].width = 70

    _sheet(wb, 'Row-by-row',
           ['status', 'reasons', 'sheet', 'cell', 'day', 'teacher (sheet)', 'teacher (DB)', 'teacher id', 'classification',
            'teacher basis', 'header as written', 'start', 'start_minute', 'duration (merged cols x 30)', 'duration by headers',
            'cell text (raw)', 'students (DB)', 'student ids', 'student evidence', 'existing lesson id', 'note / details',
            'provenance key', 'cell fill'],
           [[o.status, ', '.join(o.reasons), o.sheet, o.cell, DAY_NAMES[o.day_of_week], ws(o.teacher_raw),
             te[o.teacher_id]['full_name'].strip() if o.teacher_id else '', o.teacher_id or '', o.teacher_classification or '',
             o.teacher_basis or '', o.header_raw, hm(o.start_minute), o.start_minute, o.duration_minutes, o.duration_by_headers,
             ws(o.text_raw), names(o.student_ids), ', '.join(o.student_ids or []), o.student_basis, o.existing_lesson_id or '',
             ' | '.join([o.existing_note] + o.details if o.existing_note else o.details), o.source_key, o.fill or '']
            for o in outcomes],
           [16, 26, 10, 7, 10, 20, 24, 38, 10, 22, 9, 10, 7, 8, 8, 30, 30, 40, 22, 38, 80, 20, 10])

    _sheet(wb, 'Pending',
           ['primary reason', 'decision required', 'all reasons', 'sheet', 'cell', 'day', 'teacher (sheet)', 'start', 'duration',
            'cell text (raw)', 'details'],
           [[o.reasons[0], DECISIONS[o.reasons[0]], ', '.join(o.reasons), o.sheet, o.cell, DAY_NAMES[o.day_of_week],
             ws(o.teacher_raw), hm(o.start_minute), o.duration_minutes, ws(o.text_raw), ' | '.join(o.details)]
            for o in sorted((o for o in outcomes if o.status == 'PENDING'), key=lambda o: (list(DECISIONS).index(o.reasons[0]), o.sheet, o.row))],
           [24, 60, 30, 10, 7, 10, 20, 10, 8, 30, 100])

    _sheet(wb, 'Planned inserts',
           ['source_sheet', 'source_cell', 'teacher', 'teacher id', 'day', 'start', 'start_minute', 'duration_minutes', 'students', 'student ids'],
           [[i['source_sheet'], i['source_cell'], te[i['teacher_id']]['full_name'].strip(), i['teacher_id'], DAY_NAMES[i['day_of_week']],
             hm(i['start_minute']), i['start_minute'], i['duration_minutes'], names(i['student_ids']), ', '.join(i['student_ids'])]
            for i in items],
           [10, 18, 24, 38, 10, 8, 8, 8, 40, 80])

    _sheet(wb, 'Teacher classification',
           ['sheet spellings', 'lesson cells', 'status', 'DB teacher', 'teacher id', 'classification', 'basis', 'candidates',
            'exact lesson matches', 'evidence / decision required', 'outcomes'],
           [[' | '.join(r['sheet_spellings']), r['lesson_cells'], r['status'], r['teacher_name'] or '', r['teacher_id'] or '',
             r['classification'], r['basis'], ', '.join(r['candidates']), str(r['evidence_matches']), r['note'], str(r['outcomes'])]
            for r in classification],
           [40, 8, 12, 26, 38, 26, 30, 60, 50, 90, 50])

    _sheet(wb, 'DB lessons not in sheet',
           ['category', 'lesson id', 'teacher', 'day', 'start', 'duration', 'students', 'workbook cells (same teacher/day/student)',
            'source_sheet (10-06 import)', 'source_cell (older layout)'],
           [[l['category'], l['id'], te[l['teacher_id']]['full_name'].strip(), DAY_NAMES[l['day_of_week']], hm(l['start_minute']),
             l['duration_minutes'], names([p['student_id'] for p in l['parts']]), ', '.join(l['workbook_cells']),
             l['source_sheet'], l['source_cell']] for l in result['db_only_lessons']],
           [22, 38, 20, 10, 8, 8, 40, 40, 12, 10])

    _sheet(wb, 'Header interpretation', ['sheet', 'column', 'header as written', 'reconstructed time', 'minute'],
           [[sheet, openpyxl.utils.get_column_letter(h['col']), h['raw'], hm(h['minute']), h['minute']]
            for sheet, hs in parsed.headers.items() for h in hs], [10, 8, 10, 16, 8])

    _sheet(wb, 'Non-lesson cells', ['sheet', 'cell', 'raw value', 'merged range', 'fill', 'why not a lesson'],
           [[n['sheet'], n['cell'], repr(n['raw']) if n['raw'] is not None else '', n['merge'] or '', n['fill'] or '', n['reason']]
            for n in parsed.non_lessons], [10, 7, 20, 12, 12, 30])
    wb.save(path)


def summary_markdown(parsed, result, classification, items, plan_id, fingerprint) -> str:
    outcomes = result['outcomes']
    counts = Counter(o.status for o in outcomes)
    primary = Counter(o.reasons[0] for o in outcomes if o.status == 'PENDING')
    anyr = Counter(r for o in outcomes for r in o.reasons)
    by_class = Counter((o.teacher_classification, o.status) for o in outcomes)
    lines = [
        f'# Grid schedule import — plan `{plan_id}`', '',
        f'- Workbook SHA-256: `{parsed.sha256}`',
        f'- Booked-schedule fingerprint at plan time: `{fingerprint}`',
        f'- Source lesson cells: **{len(outcomes)}** (plus {len(parsed.non_lessons)} blank/whitespace cells that are not lessons)', '',
        '| Outcome | Cells |', '|---|---:|',
        *[f'| {k} | {counts[k]} |' for k in ('ALREADY_PRESENT', 'IMPORT', 'PENDING', 'REJECTED')],
        f'| **Total** | **{sum(counts.values())}** |', '',
        f'Planned writes: **{len(items)}** `lessons` INSERTs and **{sum(len(i["student_ids"]) for i in items)}** `lesson_participants` INSERTs. No UPDATE or DELETE.', '',
        '| Outcome by teacher classification | shift/part-time | hourly | unresolved |', '|---|---:|---:|---:|',
        *[f'| {s} | {by_class[("shift", s)]} | {by_class[("hourly", s)]} | {by_class[(None, s)]} |' for s in ('ALREADY_PRESENT', 'IMPORT', 'PENDING')],
        '', '## Pending, by primary reason', '', '| Reason | Cells (primary) | Cells (any) | Decision required |', '|---|---:|---:|---|',
        *[f'| {k} | {primary[k]} | {anyr[k]} | {DECISIONS[k]} |' for k in DECISIONS if anyr[k]],
        '', '## Active lessons from the 2026-10-06 import that this workbook does not reproduce', '',
        f'{len(result["db_only_lessons"])} lessons — kept unchanged. '
        f'{sum(1 for l in result["db_only_lessons"] if l["category"] == "CHANGED_IN_WORKBOOK")} have a workbook cell for the same teacher, day and student at another time or duration (possible reschedule); '
        f'{sum(1 for l in result["db_only_lessons"] if l["category"] == "NOT_IN_WORKBOOK")} have none.',
    ]
    if result['map_problems']:
        lines += ['', '## Teacher map problems (entries treated as UNRESOLVED)', '', *[f'- {p}' for p in result['map_problems']]]
    return '\n'.join(lines) + '\n'


def classification_markdown(classification) -> str:
    lines = ['# Teacher classification report', '',
             'Classification rule (owner, 2026-10-13): the 14 teachers configured on the schedule are the only shift-based/part-time '
             'teachers; every other reliably identified teacher is hourly. Identity is verified independently — the rule never decides who a teacher is.', '',
             '| Sheet spelling(s) | Cells | Identity | DB teacher | Teacher id | Classification | Basis | Evidence / decision required |',
             '|---|---:|---|---|---|---|---|---|']
    for r in sorted(classification, key=lambda r: (r['status'] != 'CONFIRMED', not r['on_roster'], -r['lesson_cells'])):
        lines.append(f"| {' / '.join(r['sheet_spellings'])} | {r['lesson_cells']} | {r['status']} | {r['teacher_name'] or '—'} | "
                     f"`{r['teacher_id'] or '—'}` | {r['classification']} | {r['basis'] or '—'} | {r['note']} |")
    return '\n'.join(lines) + '\n'
